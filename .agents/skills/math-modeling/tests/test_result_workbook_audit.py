from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPT = SKILL_DIR / "scripts" / "audit_result_workbook.py"


class ResultWorkbookAuditTests(unittest.TestCase):
    def make_case(self, root: Path) -> tuple[Path, Path]:
        case = root / "case"
        results = case / "results"
        results.mkdir(parents=True)
        workbook = results / "result.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Results"
        sheet["A1"] = "metric"
        sheet["B2"] = 26.0
        sheet["B2"].number_format = "0.00"
        sheet["B3"] = 0.925
        sheet["B3"].number_format = "0.0%"
        sheet["B4"] = "plan-a"
        book.create_sheet("Notes")["A1"] = "controlled fixture"
        book.save(workbook)
        book.close()
        register = {
            "schema_version": 1,
            "results": [
                {"id": "q1.objective", "value": 26.0, "absolute_tolerance": 1e-9},
                {"id": "q1.rate", "value": 0.925, "absolute_tolerance": 1e-9},
                {"id": "q1.plan", "value": "plan-a"},
            ],
        }
        (results / "result-register.json").write_text(json.dumps(register), encoding="utf-8")
        return case, workbook

    def write_plan(self, case: Path, workbook_path: Path, **updates) -> Path:
        plan = {
            "schema_version": 1,
            "workbook": "results/result.xlsx",
            "workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
            "required_sheets": ["Results", "Notes"],
            "forbid_extra_sheets": True,
            "required_cells": [
                {"id": "objective", "sheet": "Results", "cell": "B2", "expected": 26, "expected_type": "number", "absolute_tolerance": 1e-9, "required_number_format": "0.00"},
                {"id": "rate", "sheet": "Results", "cell": "B3", "expected_type": "number", "required_number_format": "0.0%"},
                {"id": "plan", "sheet": "Results", "cell": "B4", "expected_type": "string"},
            ],
            "registered_results": [
                {"result_id": "q1.objective", "sheet": "Results", "cell": "B2"},
                {"result_id": "q1.rate", "sheet": "Results", "cell": "B3"},
                {"result_id": "q1.plan", "sheet": "Results", "cell": "B4"},
            ],
        }
        plan.update(updates)
        path = case / "workbook-audit-plan.json"
        path.write_text(json.dumps(plan, indent=2) + "\n", encoding="utf-8")
        return path

    def run_audit(self, case: Path, plan: Path, name: str = "audit") -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            [sys.executable, str(SCRIPT), "--case-dir", str(case), "--plan", str(plan), "--output-dir", str(case / name), "--json"],
            capture_output=True, text=True, check=False,
        )

    def test_pass_is_hash_bound_read_only_deterministic_and_structured(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case, workbook = self.make_case(Path(temporary))
            before = workbook.read_bytes()
            plan = self.write_plan(case, workbook)
            first = self.run_audit(case, plan, "audit-a")
            second = self.run_audit(case, plan, "audit-b")
            self.assertEqual(first.returncode, 0, first.stderr)
            self.assertEqual(second.returncode, 0, second.stderr)
            self.assertEqual(workbook.read_bytes(), before)
            report = json.loads(first.stdout)
            self.assertTrue(report["ready"])
            self.assertTrue(report["input_unchanged"])
            self.assertEqual(len(report["critical_cells"]), 3)
            self.assertTrue(all(item["ok"] for item in report["registered_results"]))
            self.assertEqual(
                (case / "audit-a" / "report.json").read_bytes(),
                (case / "audit-b" / "report.json").read_bytes(),
            )

    def test_missing_sheet_blank_format_and_register_mismatch_block_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case, workbook = self.make_case(Path(temporary))
            book = load_workbook(workbook)
            del book["Notes"]
            book["Results"]["B3"] = None
            book["Results"]["B2"].number_format = "General"
            book.save(workbook)
            book.close()
            plan = self.write_plan(case, workbook)
            completed = self.run_audit(case, plan)
            self.assertEqual(completed.returncode, 2)
            report = json.loads(completed.stdout)
            joined = "\n".join(report["errors"])
            self.assertIn("required sheets are missing", joined)
            self.assertIn("critical cell is blank", joined)
            self.assertIn("number format", joined)
            self.assertIn("registered result q1.rate", joined)
            self.assertFalse((case / "audit").exists())

    def test_formula_without_cached_value_and_spreadsheet_error_are_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case, workbook = self.make_case(Path(temporary))
            book = load_workbook(workbook)
            book["Results"]["B2"] = "=13*2"
            book["Results"]["C2"] = "=#REF!"
            book.save(workbook)
            book.close()
            plan = self.write_plan(case, workbook)
            plan_data = json.loads(plan.read_text(encoding="utf-8"))
            plan_data["required_cells"][0]["allow_formula"] = True
            plan.write_text(json.dumps(plan_data), encoding="utf-8")
            completed = self.run_audit(case, plan)
            report = json.loads(completed.stdout)
            self.assertEqual(completed.returncode, 2)
            self.assertIn("formula cell Results!B2 has no cached value", report["errors"])
            self.assertTrue(any("spreadsheet error at Results!C2" in error for error in report["errors"]))

    def test_hidden_critical_locations_are_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case, workbook = self.make_case(Path(temporary))
            book = load_workbook(workbook)
            book["Results"].row_dimensions[2].hidden = True
            book["Results"].column_dimensions["B"].hidden = True
            book["Notes"].sheet_state = "hidden"
            book.save(workbook)
            book.close()
            plan = self.write_plan(case, workbook)
            plan_data = json.loads(plan.read_text(encoding="utf-8"))
            plan_data["required_cells"].append(
                {"id": "note", "sheet": "Notes", "cell": "A1", "expected_type": "string"}
            )
            plan.write_text(json.dumps(plan_data), encoding="utf-8")
            completed = self.run_audit(case, plan)
            report = json.loads(completed.stdout)
            joined = "\n".join(report["errors"])
            self.assertIn("critical cell is in a hidden row", joined)
            self.assertIn("critical cell is in a hidden column", joined)
            self.assertIn("critical cell is on a hidden sheet", joined)
            self.assertTrue(any("contains hidden rows" in warning for warning in report["warnings"]))

    def test_registered_result_can_declare_export_rounding_tolerance(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case, workbook = self.make_case(Path(temporary))
            book = load_workbook(workbook)
            book["Results"]["B2"] = 25.9999997
            book.save(workbook)
            book.close()
            plan = self.write_plan(case, workbook)
            payload = json.loads(plan.read_text(encoding="utf-8"))
            payload["required_cells"][0]["expected"] = 25.9999997
            payload["registered_results"][0]["absolute_tolerance"] = 1e-6
            plan.write_text(json.dumps(payload), encoding="utf-8")
            completed = self.run_audit(case, plan)
            self.assertEqual(completed.returncode, 0, completed.stdout)
            report = json.loads(completed.stdout)
            self.assertEqual(
                report["registered_results"][0]["absolute_tolerance"], 1e-6
            )

    def test_rejects_hash_mismatch_unsupported_extension_escape_and_existing_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case, workbook = self.make_case(Path(temporary))
            plan = self.write_plan(case, workbook, workbook_sha256="0" * 64)
            mismatch = self.run_audit(case, plan, "hash")
            self.assertEqual(mismatch.returncode, 2)
            self.assertIn("SHA-256", mismatch.stdout)
            self.assertFalse((case / "hash").exists())

            lowercase_plan = self.write_plan(case, workbook)
            lowercase_data = json.loads(lowercase_plan.read_text(encoding="utf-8"))
            lowercase_data["required_cells"][0]["cell"] = "b2"
            lowercase_plan.write_text(json.dumps(lowercase_data), encoding="utf-8")
            lowercase = self.run_audit(case, lowercase_plan, "lowercase")
            self.assertEqual(lowercase.returncode, 2)
            self.assertIn("^[A-Z]{1,3}[1-9][0-9]*$", lowercase.stdout)
            self.assertFalse((case / "lowercase").exists())

            fake = case / "results" / "legacy.xlsm"
            fake.write_bytes(workbook.read_bytes())
            unsupported = self.write_plan(case, fake, workbook="results/legacy.xlsm")
            rejected = self.run_audit(case, unsupported, "unsupported")
            self.assertEqual(rejected.returncode, 2)
            self.assertIn("only .xlsx", rejected.stdout)

            escaped_data = json.loads(unsupported.read_text(encoding="utf-8"))
            escaped_data["workbook"] = "../outside.xlsx"
            unsupported.write_text(json.dumps(escaped_data), encoding="utf-8")
            escaped = self.run_audit(case, unsupported, "escaped")
            self.assertEqual(escaped.returncode, 2)
            self.assertIn("escapes the case directory", escaped.stdout)

            existing = case / "existing"
            existing.mkdir()
            valid = self.write_plan(case, workbook)
            occupied = self.run_audit(case, valid, "existing")
            self.assertEqual(occupied.returncode, 2)
            self.assertIn("already exists", occupied.stdout)

    def test_malformed_ooxml_package_returns_structured_failure(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case, workbook = self.make_case(Path(temporary))
            with zipfile.ZipFile(workbook, "w") as package:
                package.writestr("placeholder.txt", "not a workbook")
            plan = self.write_plan(case, workbook)
            completed = self.run_audit(case, plan)
            self.assertEqual(completed.returncode, 2)
            report = json.loads(completed.stdout)
            self.assertFalse(report["ready"])
            self.assertTrue(report["errors"])
            self.assertFalse((case / "audit").exists())


if __name__ == "__main__":
    unittest.main()
