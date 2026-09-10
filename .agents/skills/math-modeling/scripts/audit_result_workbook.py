#!/usr/bin/env python
"""Audit a result XLSX without calculating, modifying, or saving it."""

from __future__ import annotations

import argparse
import json
import math
import shutil
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils.cell import coordinate_from_string

from _json_schema import load_and_validate
from _workflow_common import atomic_write_json, resolve_inside, resolve_path_inside, sha256_file
from reconcile_results import load_json, values_match


FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
SCHEMA = Path(__file__).resolve().parents[1] / "schemas" / "result-workbook-audit-plan.schema.json"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case-dir", type=Path, required=True)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--json", action="store_true", help="Emit the complete report")
    return parser.parse_args()


def _finite_scalar(value: object) -> bool:
    return not isinstance(value, float) or math.isfinite(value)


def _cell_type_matches(value: object, expected: str) -> bool:
    if expected == "boolean":
        return isinstance(value, bool)
    if expected == "number":
        return isinstance(value, (int, float)) and not isinstance(value, bool) and _finite_scalar(value)
    return isinstance(value, str)


def _is_formula_error(value: object) -> bool:
    return isinstance(value, str) and any(token in value for token in FORMULA_ERRORS)


def _load_plan(path: Path) -> dict:
    plan = load_json(path)
    load_and_validate(plan, SCHEMA, "result workbook audit plan")
    cell_ids = [item["id"] for item in plan["required_cells"]]
    if len(cell_ids) != len(set(cell_ids)):
        raise ValueError("required_cells ids must be unique")
    result_ids = [item["result_id"] for item in plan.get("registered_results", [])]
    if len(result_ids) != len(set(result_ids)):
        raise ValueError("registered_results result_id values must be unique")
    return plan


def _load_register(path: Path) -> dict[str, dict]:
    payload = load_json(path)
    if not isinstance(payload, dict) or payload.get("schema_version") != 1:
        raise ValueError("result register schema_version must be 1")
    entries = payload.get("results")
    if not isinstance(entries, list):
        raise ValueError("result register results must be an array")
    indexed: dict[str, dict] = {}
    for entry in entries:
        if not isinstance(entry, dict) or not isinstance(entry.get("id"), str):
            raise ValueError("each result register entry must have a string id")
        if entry["id"] in indexed:
            raise ValueError(f"duplicate result register id: {entry['id']}")
        if "value" not in entry or not isinstance(
            entry["value"], (str, int, float, bool)
        ):
            raise ValueError(f"result register entry {entry['id']} must have a scalar value")
        if not _finite_scalar(entry["value"]):
            raise ValueError(f"result register entry {entry['id']} must have a finite value")
        for name in ("absolute_tolerance", "relative_tolerance"):
            tolerance = entry.get(name, 0.0)
            if (
                isinstance(tolerance, bool)
                or not isinstance(tolerance, (int, float))
                or not math.isfinite(tolerance)
                or tolerance < 0
            ):
                raise ValueError(
                    f"result register entry {entry['id']} {name} must be finite and non-negative"
                )
        indexed[entry["id"]] = entry
    return indexed


def _workbook_summary(workbook) -> list[dict]:
    summaries = []
    for sheet in workbook.worksheets:
        nonempty = 0
        formulas = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    nonempty += 1
                    formulas += int(cell.data_type == "f")
        summaries.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "used_range": sheet.calculate_dimension(),
                "nonempty_cells": nonempty,
                "formula_cells": formulas,
                "merged_ranges": len(sheet.merged_cells.ranges),
                "hidden_rows": sorted(index for index, dimension in sheet.row_dimensions.items() if dimension.hidden),
                "hidden_columns": sorted(name for name, dimension in sheet.column_dimensions.items() if dimension.hidden),
            }
        )
    return summaries


def audit(case_dir: Path, plan_path: Path) -> dict:
    case_dir = case_dir.expanduser().resolve()
    plan_path = resolve_path_inside(case_dir, plan_path, "audit plan")
    report = {
        "schema_version": 1,
        "ready": False,
        "case_dir": str(case_dir),
        "plan": str(plan_path),
        "plan_sha256": None,
        "workbook": None,
        "workbook_sha256": None,
        "input_unchanged": False,
        "sheets": [],
        "critical_cells": [],
        "registered_results": [],
        "warnings": [],
        "errors": [],
        "limitations": [
            "openpyxl does not calculate formulas; formula checks use cached values saved by spreadsheet software",
            "mechanical QA does not verify column widths, clipping, pagination, print areas, charts, or visual layout",
            "opening successfully and matching registered values do not prove that the model or result is correct",
        ],
    }
    if not case_dir.is_dir():
        report["errors"].append(f"case directory does not exist: {case_dir}")
        return report
    try:
        plan = _load_plan(plan_path)
        report["plan_sha256"] = sha256_file(plan_path)
        workbook_path = resolve_inside(case_dir, plan["workbook"], "workbook")
        report["workbook"] = str(workbook_path)
        if workbook_path.suffix.lower() != ".xlsx":
            raise ValueError("only .xlsx result workbooks are supported; .xls and .xlsm require a separate reviewed path")
        if not workbook_path.is_file():
            raise ValueError(f"workbook does not exist: {workbook_path}")
        before_hash = sha256_file(workbook_path)
        report["workbook_sha256"] = before_hash
        if before_hash != plan["workbook_sha256"]:
            raise ValueError("workbook SHA-256 does not match the audit plan")
        if not zipfile.is_zipfile(workbook_path):
            raise ValueError("workbook is not a readable OOXML ZIP package")
        with zipfile.ZipFile(workbook_path) as package:
            corrupt_member = package.testzip()
        if corrupt_member is not None:
            raise ValueError(f"workbook OOXML package contains a corrupt member: {corrupt_member}")

        formula_book = None
        value_book = None
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            try:
                formula_book = load_workbook(
                    workbook_path, data_only=False, read_only=False, keep_links=True
                )
                value_book = load_workbook(
                    workbook_path, data_only=True, read_only=False, keep_links=True
                )
            except Exception as exc:
                if formula_book is not None:
                    formula_book.close()
                raise ValueError(f"cannot parse workbook {workbook_path}: {exc}") from exc
        report["warnings"].extend(f"openpyxl: {item.message}" for item in caught)
        try:
            report["sheets"] = _workbook_summary(formula_book)
            names = formula_book.sheetnames
            missing = [name for name in plan["required_sheets"] if name not in names]
            if missing:
                report["errors"].append(f"required sheets are missing: {', '.join(missing)}")
            if plan.get("forbid_extra_sheets", False):
                extras = [name for name in names if name not in plan["required_sheets"]]
                if extras:
                    report["errors"].append(f"unexpected sheets are present: {', '.join(extras)}")
            if getattr(formula_book, "_external_links", []):
                report["warnings"].append("workbook contains external links; review them manually")

            for summary in report["sheets"]:
                if summary["state"] != "visible":
                    report["warnings"].append(f"sheet {summary['name']} is {summary['state']}")
                if summary["hidden_rows"]:
                    report["warnings"].append(f"sheet {summary['name']} contains hidden rows")
                if summary["hidden_columns"]:
                    report["warnings"].append(f"sheet {summary['name']} contains hidden columns")
                if summary["merged_ranges"]:
                    report["warnings"].append(f"sheet {summary['name']} contains merged cells")

            for sheet in formula_book.worksheets:
                values_sheet = value_book[sheet.title]
                for row in sheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        cached = values_sheet[cell.coordinate].value
                        if cell.data_type == "f" and cached is None:
                            report["errors"].append(f"formula cell {sheet.title}!{cell.coordinate} has no cached value")
                        cached_cell = values_sheet[cell.coordinate]
                        broken_reference = cell.data_type == "f" and "#REF!" in str(value)
                        if (
                            broken_reference
                            or cell.data_type == "e"
                            or cached_cell.data_type == "e"
                            or (cell.data_type != "f" and _is_formula_error(value))
                            or _is_formula_error(cached)
                        ):
                            report["errors"].append(f"spreadsheet error at {sheet.title}!{cell.coordinate}: {value!r}")

            for check in plan["required_cells"]:
                item = {"id": check["id"], "sheet": check["sheet"], "cell": check["cell"], "ok": False, "errors": []}
                report["critical_cells"].append(item)
                if check["sheet"] not in formula_book.sheetnames:
                    item["errors"].append("sheet does not exist")
                    continue
                sheet = formula_book[check["sheet"]]
                value_sheet = value_book[check["sheet"]]
                cell = sheet[check["cell"]]
                actual = value_sheet[check["cell"]].value if cell.data_type == "f" else cell.value
                item.update({"value": actual, "formula": cell.value if cell.data_type == "f" else None, "number_format": cell.number_format})
                _, row_index = coordinate_from_string(check["cell"] )
                if sheet.sheet_state != "visible":
                    item["errors"].append("critical cell is on a hidden sheet")
                if sheet.row_dimensions[row_index].hidden:
                    item["errors"].append("critical cell is in a hidden row")
                column_name, _ = coordinate_from_string(check["cell"] )
                if sheet.column_dimensions[column_name].hidden:
                    item["errors"].append("critical cell is in a hidden column")
                if cell.data_type == "f" and not check.get("allow_formula", False):
                    item["errors"].append("formula is not allowed for this critical cell")
                if actual is None and not check.get("allow_blank", False):
                    item["errors"].append("critical cell is blank")
                if not _finite_scalar(actual):
                    item["errors"].append("critical numeric value must be finite")
                if "expected_type" in check and actual is not None and not _cell_type_matches(actual, check["expected_type"]):
                    item["errors"].append(f"value is not of expected type {check['expected_type']}")
                if "expected" in check and not values_match(check["expected"], actual, check.get("absolute_tolerance", 0.0), check.get("relative_tolerance", 0.0)):
                    item["errors"].append("value does not match expected value")
                if "required_number_format" in check and cell.number_format != check["required_number_format"]:
                    item["errors"].append(f"number format {cell.number_format!r} does not match {check['required_number_format']!r}")
                item["ok"] = not item["errors"]

            mappings = plan.get("registered_results", [])
            register = {}
            if mappings:
                register_path = resolve_inside(case_dir, plan.get("result_register", "results/result-register.json"), "result register")
                register = _load_register(register_path)
                report["result_register"] = str(register_path)
                report["result_register_sha256"] = sha256_file(register_path)
            for mapping in mappings:
                item = {**mapping, "ok": False, "errors": []}
                report["registered_results"].append(item)
                entry = register.get(mapping["result_id"])
                if entry is None:
                    item["errors"].append("result id is missing from the register")
                elif mapping["sheet"] not in value_book.sheetnames:
                    item["errors"].append("sheet does not exist")
                else:
                    actual = value_book[mapping["sheet"]][mapping["cell"]].value
                    item["registered_value"] = entry.get("value")
                    item["workbook_value"] = actual
                    absolute = mapping.get(
                        "absolute_tolerance", entry.get("absolute_tolerance", 0.0)
                    )
                    relative = mapping.get(
                        "relative_tolerance", entry.get("relative_tolerance", 0.0)
                    )
                    item["absolute_tolerance"] = absolute
                    item["relative_tolerance"] = relative
                    if not values_match(entry.get("value"), actual, absolute, relative):
                        item["errors"].append("workbook value does not match the registered result")
                item["ok"] = not item["errors"]
        finally:
            formula_book.close()
            value_book.close()

        for item in report["critical_cells"]:
            report["errors"].extend(f"critical cell {item['id']}: {error}" for error in item["errors"])
        for item in report["registered_results"]:
            report["errors"].extend(f"registered result {item['result_id']}: {error}" for error in item["errors"])
        after_hash = sha256_file(workbook_path)
        report["input_unchanged"] = after_hash == before_hash
        if not report["input_unchanged"]:
            report["errors"].append("workbook changed during the audit")
    except (
        OSError,
        ValueError,
        KeyError,
        TypeError,
        InvalidFileException,
        zipfile.BadZipFile,
    ) as exc:
        report["errors"].append(str(exc))
    report["ready"] = not report["errors"]
    return report


def commit_report(output_dir: Path, report: dict) -> None:
    if output_dir.exists():
        raise ValueError(f"output directory already exists: {output_dir}")
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{output_dir.name}.", dir=output_dir.parent))
    try:
        atomic_write_json(staging / "report.json", report)
        staging.replace(output_dir)
        staging = None
    finally:
        if staging is not None:
            shutil.rmtree(staging, ignore_errors=True)


def main() -> int:
    args = parse_args()
    case_dir = args.case_dir.expanduser().resolve()
    try:
        output_dir = resolve_path_inside(case_dir, args.output_dir, "output directory")
        if output_dir.exists():
            raise ValueError(f"output directory already exists: {output_dir}")
        report = audit(case_dir, args.plan)
        if report["ready"]:
            commit_report(output_dir, report)
    except (OSError, ValueError) as exc:
        report = {"schema_version": 1, "ready": False, "errors": [str(exc)], "warnings": []}
    if args.json:
        json.dump(report, sys.stdout, ensure_ascii=False, indent=2)
        print()
    else:
        print(f"ready={str(report['ready']).lower()}")
        for warning in report.get("warnings", []):
            print(f"warning: {warning}")
        for error in report.get("errors", []):
            print(f"error: {error}")
    return 0 if report["ready"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
