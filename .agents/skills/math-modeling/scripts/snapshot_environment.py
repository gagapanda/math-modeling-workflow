#!/usr/bin/env python
"""Create a reproducible modeling-environment snapshot and optional smoke report."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import os
import platform
import re
import string
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from xml.dom import Node, minidom
from xml.parsers.expat import ExpatError

import preflight


SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_REQUIREMENTS = SKILL_DIR.parents[2] / "requirements-modeling.txt"
EXPECTED_MATLAB_SMOKE_VALUE = 385
XMP_NAMESPACE = "http://ns.adobe.com/xap/1.0/"
DC_NAMESPACE = "http://purl.org/dc/elements/1.1/"
RDF_NAMESPACE = "http://www.w3.org/1999/02/22-rdf-syntax-ns#"
XMLNS_NAMESPACE = "http://www.w3.org/2000/xmlns/"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=SKILL_DIR.parents[2])
    parser.add_argument("--requirements", type=Path, default=DEFAULT_REQUIREMENTS)
    parser.add_argument("--output", type=Path, help="Write the JSON report atomically")
    parser.add_argument(
        "--generated-at",
        help=(
            "Explicit timezone-aware audit timestamp; required with --active-smoke "
            "and applied to generated PDF metadata"
        ),
    )
    parser.add_argument(
        "--active-smoke",
        action="store_true",
        help="Run controlled DOCX/PDF/render and XeLaTeX smoke chains",
    )
    parser.add_argument(
        "--smoke-dir",
        type=Path,
        help="Workspace for active smoke artifacts; required with --active-smoke",
    )
    parser.add_argument(
        "--matlab-mcp-evidence",
        type=Path,
        help="Validated evidence from a genuine MATLAB MCP call",
    )
    parser.add_argument("--json", action="store_true", help="Emit JSON to stdout")
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_audit_timestamp(value: str | None) -> tuple[datetime, str]:
    if value is None:
        parsed = datetime.now().astimezone()
    else:
        try:
            parsed = datetime.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("generated_at must be an ISO 8601 timestamp") from exc
        if parsed.tzinfo is None or parsed.utcoffset() is None:
            raise ValueError("generated_at must include a timezone offset")
    return parsed, parsed.isoformat(timespec="seconds")


def pdf_date_string(timestamp: datetime) -> str:
    offset = timestamp.utcoffset()
    if offset is None:
        raise ValueError("PDF metadata timestamp must include a timezone offset")
    offset_seconds = int(offset.total_seconds())
    if offset_seconds % 60:
        raise ValueError("PDF metadata timezone offset must use whole minutes")
    if offset_seconds == 0:
        suffix = "Z"
    else:
        sign = "+" if offset_seconds > 0 else "-"
        hours, minutes = divmod(abs(offset_seconds) // 60, 60)
        suffix = f"{sign}{hours:02d}'{minutes:02d}'"
    return timestamp.strftime("D:%Y%m%d%H%M%S") + suffix


def parse_xmp_document(data: bytes) -> minidom.Document:
    try:
        return minidom.parseString(data)
    except ExpatError as exc:
        raise ValueError(f"existing PDF XMP is invalid XML: {exc}") from exc


def replace_element_text(element: minidom.Element, value: str) -> None:
    for child in list(element.childNodes):
        element.removeChild(child)
    element.appendChild(element.ownerDocument.createTextNode(value))


def xmp_field_values(document: minidom.Document, local_name: str) -> list[str]:
    values = []
    descriptions = document.getElementsByTagNameNS(RDF_NAMESPACE, "Description")
    for description in descriptions:
        attribute = description.getAttributeNodeNS(XMP_NAMESPACE, local_name)
        if attribute is not None:
            values.append(attribute.value)
    for element in document.getElementsByTagNameNS(XMP_NAMESPACE, local_name):
        values.append(
            "".join(
                child.data
                for child in element.childNodes
                if child.nodeType in {Node.TEXT_NODE, Node.CDATA_SECTION_NODE}
            ).strip()
        )
    return values


def dc_date_values(document: minidom.Document) -> list[str]:
    values = []
    for date_element in document.getElementsByTagNameNS(DC_NAMESPACE, "date"):
        for item in date_element.getElementsByTagNameNS(RDF_NAMESPACE, "li"):
            values.append(
                "".join(
                    child.data
                    for child in item.childNodes
                    if child.nodeType in {Node.TEXT_NODE, Node.CDATA_SECTION_NODE}
                ).strip()
            )
    return values


def normalize_xmp_bytes(data: bytes, audit_time: datetime) -> bytes:
    document = parse_xmp_document(data)
    audit_iso = audit_time.isoformat(timespec="seconds")
    descriptions = list(
        document.getElementsByTagNameNS(RDF_NAMESPACE, "Description")
    )
    if not descriptions:
        raise ValueError("existing PDF XMP has no rdf:Description element")

    for local_name in ("CreateDate", "ModifyDate", "MetadataDate"):
        found = False
        for description in descriptions:
            attribute = description.getAttributeNodeNS(XMP_NAMESPACE, local_name)
            if attribute is not None:
                attribute.value = audit_iso
                found = True
        for element in document.getElementsByTagNameNS(XMP_NAMESPACE, local_name):
            replace_element_text(element, audit_iso)
            found = True
        if not found:
            target = descriptions[-1]
            target.setAttributeNS(XMLNS_NAMESPACE, "xmlns:xmp", XMP_NAMESPACE)
            element = document.createElementNS(XMP_NAMESPACE, f"xmp:{local_name}")
            element.appendChild(document.createTextNode(audit_iso))
            target.appendChild(element)

    for date_element in document.getElementsByTagNameNS(DC_NAMESPACE, "date"):
        for item in date_element.getElementsByTagNameNS(RDF_NAMESPACE, "li"):
            replace_element_text(item, audit_iso)

    return document.toxml(encoding="utf-8")


def raw_xmp_bytes(reader) -> bytes | None:
    metadata_reference = reader.root_object.get("/Metadata")
    if metadata_reference is None:
        return None
    metadata_stream = metadata_reference.get_object()
    return metadata_stream.get_data()


def normalize_pdf_metadata(path: Path, audit_time: datetime) -> dict:
    from pypdf import PdfReader, PdfWriter

    resolved = path.expanduser().resolve()
    if not resolved.is_file():
        raise ValueError(f"PDF does not exist: {resolved}")
    pdf_date = pdf_date_string(audit_time)
    reader = PdfReader(resolved)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    metadata = {
        str(key): str(value)
        for key, value in (reader.metadata or {}).items()
        if value is not None
    }
    metadata["/CreationDate"] = pdf_date
    metadata["/ModDate"] = pdf_date
    writer.add_metadata(metadata)
    original_xmp = raw_xmp_bytes(reader)
    xmp_present = original_xmp is not None
    if original_xmp is not None:
        writer.xmp_metadata = normalize_xmp_bytes(original_xmp, audit_time)

    expected_iso = audit_time.isoformat(timespec="seconds")
    temporary = resolved.with_name(f".{resolved.name}.{os.urandom(6).hex()}.tmp")
    try:
        with temporary.open("wb") as stream:
            writer.write(stream)

        verified_reader = PdfReader(temporary)
        verified_metadata = verified_reader.metadata or {}
        creation_date = str(verified_metadata.get("/CreationDate", ""))
        modification_date = str(verified_metadata.get("/ModDate", ""))
        verified_xmp_bytes = raw_xmp_bytes(verified_reader)
        xmp_values = {
            "present": verified_xmp_bytes is not None,
            "create_date": None,
            "modify_date": None,
            "metadata_date": None,
            "dc_dates": [],
        }
        xmp_verified = not xmp_present
        if verified_xmp_bytes is not None:
            verified_document = parse_xmp_document(verified_xmp_bytes)
            field_values = {
                "create_date": xmp_field_values(verified_document, "CreateDate"),
                "modify_date": xmp_field_values(verified_document, "ModifyDate"),
                "metadata_date": xmp_field_values(verified_document, "MetadataDate"),
            }
            dates = dc_date_values(verified_document)
            xmp_values.update(
                {
                    key: values[0] if values else None
                    for key, values in field_values.items()
                }
            )
            xmp_values["dc_dates"] = dates
            xmp_verified = all(
                values and all(value == expected_iso for value in values)
                for values in field_values.values()
            ) and all(value == expected_iso for value in dates)
            try:
                verified_reader.xmp_metadata
            except Exception as exc:
                raise ValueError(f"normalized PDF XMP cannot be reopened: {exc}") from exc

        verified = (
            creation_date == pdf_date
            and modification_date == pdf_date
            and xmp_values["present"] == xmp_present
            and xmp_verified
        )
        if not verified:
            raise ValueError(f"PDF metadata verification failed: {resolved}")
        temporary.replace(resolved)
    finally:
        temporary.unlink(missing_ok=True)

    return {
        "applied": True,
        "verified": True,
        "audit_time": expected_iso,
        "pdf_date": pdf_date,
        "creation_date": creation_date,
        "modification_date": modification_date,
        "xmp": xmp_values,
    }


def run_command(
    command: list[str], cwd: Path, timeout: int = 60, tail_lines: int | None = 20
) -> dict:
    try:
        completed = subprocess.run(
            command,
            cwd=cwd,
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        return {"ok": False, "exit_code": None, "output": str(exc)}
    output = "\n".join(
        part.strip() for part in (completed.stdout, completed.stderr) if part.strip()
    )
    if tail_lines is not None:
        output = "\n".join(output.splitlines()[-tail_lines:])
    return {
        "ok": completed.returncode == 0,
        "exit_code": completed.returncode,
        "output": output,
    }


def parse_pinned_requirements(path: Path) -> dict[str, str]:
    pins: dict[str, str] = {}
    for line_number, raw in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = raw.split("#", 1)[0].strip()
        if not line:
            continue
        match = re.fullmatch(r"([A-Za-z0-9_.-]+)==([^\s]+)", line)
        if match is None:
            raise ValueError(
                f"{path}:{line_number}: core requirements must use exact == pins"
            )
        name, version = match.groups()
        normalized = re.sub(r"[-_.]+", "-", name).casefold()
        if normalized in pins:
            raise ValueError(f"{path}:{line_number}: duplicate requirement {name}")
        pins[normalized] = version
    if not pins:
        raise ValueError(f"{path}: no pinned requirements found")
    return pins


def dependency_report(requirements: Path) -> dict:
    expected = parse_pinned_requirements(requirements)
    packages = []
    mismatches = []
    for name, wanted in sorted(expected.items()):
        try:
            installed = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            installed = None
        matches = installed == wanted
        packages.append(
            {"name": name, "expected": wanted, "installed": installed, "matches": matches}
        )
        if not matches:
            mismatches.append(name)
    pip_check = run_command(
        [sys.executable, "-m", "pip", "check"], requirements.parent, timeout=120
    )
    pip_freeze = run_command(
        [sys.executable, "-m", "pip", "freeze", "--all"],
        requirements.parent,
        timeout=120,
        tail_lines=None,
    )
    freeze_lines = (
        sorted(line for line in pip_freeze["output"].splitlines() if line.strip())
        if pip_freeze["ok"]
        else []
    )
    return {
        "requirements_path": str(requirements.resolve()),
        "requirements_sha256": sha256_file(requirements),
        "core_packages": packages,
        "core_matches": not mismatches,
        "mismatches": mismatches,
        "pip_check": pip_check,
        "pip_freeze": freeze_lines,
        "pip_freeze_sha256": hashlib.sha256(
            ("\n".join(freeze_lines) + "\n").encode("utf-8")
        ).hexdigest(),
    }


def version_probe(path: str | None, *arguments: str) -> dict:
    if path is None:
        return {"path": None, "available": False, "usable": False, "version": None}
    result = run_command([path, *arguments], Path.cwd())
    lines = [line.strip() for line in result["output"].splitlines() if line.strip()]
    version = next(
        (line for line in lines if re.search(r"\d+(?:\.\d+)+|TeX Live \d{4}", line)),
        lines[-1] if lines else None,
    )
    return {
        "path": str(Path(path).resolve()),
        "available": True,
        "usable": result["ok"],
        "version": version,
        "probe_exit_code": result["exit_code"],
    }


def find_tesseract() -> str | None:
    discovered = preflight.find_executable("tesseract")
    if discovered:
        return discovered
    if os.name == "nt":
        candidate = Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Tesseract-OCR" / "tesseract.exe"
        if candidate.is_file():
            return str(candidate)
    return None


def windows_fixed_drive_roots() -> list[Path]:
    if os.name != "nt":
        return []
    return [Path(f"{letter}:\\") for letter in string.ascii_uppercase if Path(f"{letter}:\\").is_dir()]


def find_tex_executable(name: str) -> str | None:
    discovered = preflight.find_executable(name)
    if discovered:
        return discovered
    if os.name != "nt":
        return None

    executable = f"{name}.exe"
    candidates = [
        Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
        / "MiKTeX"
        / "miktex"
        / "bin"
        / "x64"
        / executable,
        Path(os.environ.get("LOCALAPPDATA", ""))
        / "Programs"
        / "MiKTeX"
        / "miktex"
        / "bin"
        / "x64"
        / executable,
        Path(os.environ.get("APPDATA", ""))
        / "TinyTeX"
        / "bin"
        / "windows"
        / executable,
    ]
    for drive in windows_fixed_drive_roots():
        for parent in (drive / "texlive", drive / "software" / "texlive"):
            if not parent.is_dir():
                continue
            releases = sorted(
                (path for path in parent.iterdir() if path.is_dir()),
                key=lambda path: path.name,
                reverse=True,
            )
            candidates.extend(
                release / "bin" / "windows" / executable for release in releases
            )
    for candidate in dict.fromkeys(candidates):
        if candidate.is_file():
            return str(candidate.resolve())
    return None


def discover_tools() -> tuple[dict, dict]:
    process_names, process_error = preflight.running_process_names()
    poppler = {
        name: preflight.resolve_poppler_tool(name)
        for name in ("pdftoppm", "pdfinfo", "pdftocairo")
    }
    paths = {
        "word": preflight.find_word(),
        "libreoffice": preflight.find_libreoffice(),
        "pdftoppm": poppler["pdftoppm"]["path"],
        "pdfinfo": poppler["pdfinfo"]["path"],
        "pdftocairo": poppler["pdftocairo"]["path"],
        "xelatex": find_tex_executable("xelatex"),
        "latexmk": find_tex_executable("latexmk"),
        "matlab": preflight.find_executable("matlab"),
        "tesseract": find_tesseract(),
    }
    tools = {
        "word": {
            "path": paths["word"],
            "available": paths["word"] is not None,
            "usable": preflight.probe_word_com(paths["word"])[0],
            "version": None,
        },
        "libreoffice": version_probe(paths["libreoffice"], "--version"),
        "pdftoppm": {**version_probe(paths["pdftoppm"], "-v"), "resolution": poppler["pdftoppm"]},
        "pdfinfo": {**version_probe(paths["pdfinfo"], "-v"), "resolution": poppler["pdfinfo"]},
        "pdftocairo": {**version_probe(paths["pdftocairo"], "-v"), "resolution": poppler["pdftocairo"]},
        "xelatex": version_probe(paths["xelatex"], "--version"),
        "latexmk": version_probe(paths["latexmk"], "-v"),
        "matlab": {
            "path": paths["matlab"],
            "available": paths["matlab"] is not None,
            "usable": paths["matlab"] is not None,
            "version": None,
        },
        "tesseract": version_probe(paths["tesseract"], "--version"),
    }
    matlab_batch = preflight.assess_matlab_batch(
        paths["matlab"], process_names, process_error
    )
    return tools, matlab_batch


def load_matlab_mcp_evidence(path: Path | None) -> dict:
    if path is None:
        return {"provided": False, "valid": False, "reason": "evidence not provided"}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return {"provided": True, "valid": False, "reason": str(exc)}
    required = {
        "schema_version": 1,
        "source": "matlab-mcp",
        "expected_value": EXPECTED_MATLAB_SMOKE_VALUE,
        "observed_value": EXPECTED_MATLAB_SMOKE_VALUE,
    }
    errors = [f"{key} must be {value!r}" for key, value in required.items() if payload.get(key) != value]
    for key in ("verified_at", "version", "release", "smoke_expression"):
        if not isinstance(payload.get(key), str) or not payload[key].strip():
            errors.append(f"{key} must be a nonempty string")
    return {
        "provided": True,
        "valid": not errors,
        "path": str(path.resolve()),
        "sha256": sha256_file(path),
        "errors": errors,
        "evidence": payload,
    }


def pdf_page_size(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(path)
    first = reader.pages[0]
    width = float(first.mediabox.width)
    height = float(first.mediabox.height)
    return {
        "page_count": len(reader.pages),
        "width_points": round(width, 2),
        "height_points": round(height, 2),
        "a4": abs(width - 595.28) <= 3 and abs(height - 841.89) <= 3,
    }


def image_is_nonblank(path: Path) -> bool:
    from PIL import Image, ImageStat

    with Image.open(path) as image:
        extrema = ImageStat.Stat(image.convert("RGB")).extrema
    return any(low != high for low, high in extrema)


def smoke_docx_pdf(tools: dict, root: Path, audit_time: datetime) -> dict:
    from docx import Document
    from docx.shared import Mm

    libreoffice = tools["libreoffice"]["path"]
    renderer = tools["pdftoppm"]["path"] or tools["pdftocairo"]["path"]
    pdfinfo = tools["pdfinfo"]["path"]
    if not libreoffice or not renderer or not pdfinfo:
        return {"passed": False, "error": "LibreOffice, Poppler renderer, and pdfinfo are required"}
    root.mkdir(parents=True, exist_ok=True)
    docx_path = root / "environment-smoke.docx"
    pdf_path = root / "environment-smoke.pdf"
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    document.add_heading("Modeling environment smoke test", level=1)
    document.add_paragraph("Deterministic DOCX to PDF to PNG verification.")
    document.save(docx_path)
    with tempfile.TemporaryDirectory(prefix="lo-profile-", dir=root) as profile:
        conversion = run_command(
            [
                libreoffice,
                f"-env:UserInstallation={Path(profile).resolve().as_uri()}",
                "--headless",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                str(root),
                str(docx_path),
            ],
            root,
            timeout=180,
        )
    if not conversion["ok"] or not pdf_path.is_file():
        return {"passed": False, "conversion": conversion, "error": "DOCX conversion failed"}
    try:
        normalized_metadata = normalize_pdf_metadata(pdf_path, audit_time)
    except (OSError, ValueError) as exc:
        return {
            "passed": False,
            "conversion": conversion,
            "error": f"PDF metadata normalization failed: {exc}",
        }
    info = run_command([pdfinfo, str(pdf_path)], root)
    prefix = root / "environment-smoke-page"
    render_command = [renderer, "-png", "-r", "96", str(pdf_path), str(prefix)]
    rendering = run_command(render_command, root, timeout=180)
    pages = sorted(root.glob("environment-smoke-page-*.png"))
    nonblank = bool(pages) and all(image_is_nonblank(page) for page in pages)
    size = pdf_page_size(pdf_path)
    passed = conversion["ok"] and info["ok"] and rendering["ok"] and size["a4"] and nonblank
    return {
        "passed": passed,
        "docx": str(docx_path.resolve()),
        "pdf": str(pdf_path.resolve()),
        "rendered_pages": [str(page.resolve()) for page in pages],
        "pdf_size": size,
        "nonblank": nonblank,
        "conversion": conversion,
        "metadata_normalization": normalized_metadata,
        "pdfinfo": info,
        "rendering": rendering,
    }


def smoke_xelatex(tools: dict, root: Path, audit_time: datetime) -> dict:
    latexmk = tools["latexmk"]["path"]
    xelatex = tools["xelatex"]["path"]
    if not latexmk or not xelatex:
        return {"passed": False, "error": "latexmk and XeLaTeX are required"}
    root.mkdir(parents=True, exist_ok=True)
    source = root / "environment-smoke.tex"
    source.write_text(
        "\\documentclass[a4paper]{article}\n"
        "\\usepackage[margin=25mm]{geometry}\n"
        "\\begin{document}\n"
        "Modeling environment XeLaTeX smoke test.\\par\n"
        "$\\sum_{i=1}^{10} i^2 = 385$.\n"
        "\\end{document}\n",
        encoding="ascii",
    )
    build = run_command(
        [
            latexmk,
            "-xelatex",
            "-interaction=nonstopmode",
            "-halt-on-error",
            f"-outdir={root}",
            str(source),
        ],
        root,
        timeout=180,
    )
    pdf = root / "environment-smoke.pdf"
    normalized_metadata = None
    normalization_error = None
    if pdf.is_file():
        try:
            normalized_metadata = normalize_pdf_metadata(pdf, audit_time)
        except (OSError, ValueError) as exc:
            normalization_error = str(exc)
    size = pdf_page_size(pdf) if pdf.is_file() else None
    return {
        "passed": bool(
            build["ok"]
            and normalized_metadata
            and normalized_metadata["verified"]
            and size
            and size["a4"]
        ),
        "source": str(source.resolve()),
        "pdf": str(pdf.resolve()) if pdf.is_file() else None,
        "pdf_size": size,
        "metadata_normalization": normalized_metadata,
        "metadata_normalization_error": normalization_error,
        "build": build,
    }


def run_active_smoke(tools: dict, smoke_dir: Path, audit_time: datetime) -> dict:
    resolved = smoke_dir.expanduser().resolve()
    resolved.mkdir(parents=True, exist_ok=True)
    docx = smoke_docx_pdf(tools, resolved / "docx-pdf-render", audit_time)
    latex = smoke_xelatex(tools, resolved / "xelatex", audit_time)
    return {
        "requested": True,
        "directory": str(resolved),
        "audit_time": audit_time.isoformat(timespec="seconds"),
        "passed": docx["passed"] and latex["passed"],
        "docx_pdf_render": docx,
        "xelatex": latex,
        "matlab_batch_started": False,
    }


def build_report(args: argparse.Namespace) -> dict:
    project_root = args.project_root.expanduser().resolve()
    requirements = args.requirements.expanduser().resolve()
    if not project_root.is_dir():
        raise ValueError(f"project root does not exist: {project_root}")
    if not requirements.is_file():
        raise ValueError(f"requirements file does not exist: {requirements}")
    if args.active_smoke and args.smoke_dir is None:
        raise ValueError("--smoke-dir is required with --active-smoke")
    generated_at_argument = getattr(args, "generated_at", None)
    if args.active_smoke and generated_at_argument is None:
        raise ValueError("--generated-at is required with --active-smoke")
    audit_time, generated_at = parse_audit_timestamp(generated_at_argument)
    dependencies = dependency_report(requirements)
    tools, matlab_batch = discover_tools()
    matlab_mcp = load_matlab_mcp_evidence(args.matlab_mcp_evidence)
    smoke = (
        run_active_smoke(tools, args.smoke_dir, audit_time)
        if args.active_smoke
        else {"requested": False, "passed": None, "matlab_batch_started": False}
    )
    wheelhouse = project_root / "wheelhouse"
    wheel_count = len(list(wheelhouse.glob("*.whl"))) if wheelhouse.is_dir() else 0
    tesseract_recovery = project_root / ".skill-audit" / "tool-cache" / "tesseract" / "recovery-manifest.json"
    warnings = []
    if tools["word"]["available"] and not tools["word"]["usable"]:
        warnings.append("Microsoft Word is registered but COM startup is unavailable; use LibreOffice")
    if not matlab_batch["safe_to_start"]:
        warnings.append("independent MATLAB batch startup is not safe; use validated MATLAB MCP evidence")
    if wheel_count == 0:
        warnings.append("no validated local wheelhouse exists; fully offline Python recreation is not proven")
    errors = []
    if not dependencies["core_matches"] or not dependencies["pip_check"]["ok"]:
        errors.append("Python dependency baseline is not healthy")
    if args.active_smoke and not smoke["passed"]:
        errors.append("one or more active tool-chain smoke checks failed")
    if args.matlab_mcp_evidence and not matlab_mcp["valid"]:
        errors.append("MATLAB MCP evidence is invalid")
    return {
        "schema_version": 1,
        "command": "snapshot_environment",
        "generated_at": generated_at,
        "project_root": str(project_root),
        "ready": not errors,
        "platform": {
            "system": platform.system(),
            "release": platform.release(),
            "version": platform.version(),
            "machine": platform.machine(),
        },
        "python": {
            "executable": str(Path(sys.executable).resolve()),
            "version": platform.python_version(),
            "implementation": platform.python_implementation(),
            "architecture": platform.architecture()[0],
        },
        "dependencies": dependencies,
        "tools": tools,
        "matlab_batch": matlab_batch,
        "matlab_mcp": matlab_mcp,
        "fallback_order": {
            "docx_to_pdf": ["Microsoft Word COM when usable", "LibreOffice isolated profile"],
            "pdf_to_images": ["actual pdftoppm executable", "pdftocairo"],
            "matlab": ["MATLAB MCP with hash-bound evidence", "independent batch after safe startup probe"],
            "paper": ["controlled XeLaTeX template", "enhanced Word export path"],
            "ocr": ["Tesseract chi_sim for benchmark-approved prose retrieval", "manual visual formula review"],
        },
        "offline_recovery": {
            "core_requirements_locked": True,
            "wheelhouse_path": str(wheelhouse.resolve()),
            "wheel_count": wheel_count,
            "fully_offline_python_recreation_verified": False,
            "tesseract_recovery_manifest": {
                "path": str(tesseract_recovery.resolve()),
                "available": tesseract_recovery.is_file(),
                "sha256": sha256_file(tesseract_recovery) if tesseract_recovery.is_file() else None,
            },
        },
        "smoke": smoke,
        "warnings": warnings,
        "errors": errors,
    }


def write_json_atomic(path: Path, payload: dict) -> None:
    resolved = path.expanduser().resolve()
    resolved.parent.mkdir(parents=True, exist_ok=True)
    temporary = resolved.with_name(f".{resolved.name}.{os.urandom(6).hex()}.tmp")
    try:
        temporary.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        temporary.replace(resolved)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> int:
    args = parse_args()
    try:
        report = build_report(args)
    except (OSError, ValueError, subprocess.SubprocessError) as exc:
        report = {
            "schema_version": 1,
            "command": "snapshot_environment",
            "ready": False,
            "errors": [str(exc)],
            "warnings": [],
        }
    if args.output:
        write_json_atomic(args.output, report)
    if args.json or not args.output:
        json.dump(report, sys.stdout, ensure_ascii=False, indent=2)
        print()
    return 0 if report["ready"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
