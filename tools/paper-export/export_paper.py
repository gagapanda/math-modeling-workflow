#!/usr/bin/env python
"""Local Markdown -> native-equation DOCX -> PDF -> all-page PNG evidence.

New orchestration, not a copy of the legacy exporter. No network assets, macros,
Word automation, shell commands, bundled binaries or automatic review acceptance.
"""
from __future__ import annotations

import argparse
import base64
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from urllib.parse import unquote, urlsplit
import uuid

READER = "markdown-raw_html-raw_tex-yaml_metadata_block-pandoc_title_block"
TAG = re.compile(r"\\tag\{([0-9]+(?:[.\-][0-9]+)*)\}\s*$")


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def inside(root, path):
    root, path = Path(root).resolve(), Path(path).resolve()
    if not path.is_relative_to(root):
        raise ValueError("path escapes asset root")
    return path


def executable(value):
    found = shutil.which(value)
    if not found or Path(found).suffix.lower() in {".bat", ".cmd", ".ps1"}:
        raise ValueError(f"missing native executable: {value}")
    return str(Path(found).resolve())


def run(command, *, timeout, cwd=None, data=None):
    """Bounded process execution; timeout terminates only this process tree."""
    kwargs = {"start_new_session": True} if os.name != "nt" else {}
    process = subprocess.Popen(command, cwd=cwd, stdin=subprocess.PIPE,
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE, **kwargs)
    try:
        stdout, stderr = process.communicate(data, timeout=timeout)
    except subprocess.TimeoutExpired:
        if os.name == "nt":
            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                           capture_output=True, timeout=15, check=False)
        else:
            import signal
            os.killpg(process.pid, signal.SIGKILL)
        process.kill()
        process.communicate()
        raise RuntimeError(f"tool timed out: {Path(command[0]).name}")
    if process.returncode:
        raise RuntimeError(f"{Path(command[0]).name} failed ({process.returncode}): "
                           + stderr.decode("utf-8", errors="replace")[-4000:])
    return stdout


def prepare_ast(ast, source, asset_root):
    """Validate document content before passing a JSON AST to the writer."""
    from PIL import Image
    ast = deepcopy(ast)
    ast["meta"] = {}  # no author/title/remote template metadata
    labels, assets = [], {}
    math_count = 0

    def visit(node):
        nonlocal math_count
        if isinstance(node, dict):
            kind, value = node.get("t"), node.get("c")
            if kind in {"RawBlock", "RawInline"}:
                raise ValueError("raw document instructions are not supported")
            if kind == "Image":
                url = value[2][0]
                parsed = urlsplit(url)
                if parsed.scheme or parsed.netloc or parsed.query or parsed.fragment or "\\" in url:
                    raise ValueError("only local PNG/JPEG assets are supported")
                path = inside(asset_root, Path(source).parent / unquote(parsed.path))
                if not path.is_file() or path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
                    raise ValueError("image must be an existing PNG/JPEG")
                if path.stat().st_size > 20_000_000:
                    raise ValueError("image exceeds 20 MB limit")
                with Image.open(path) as image:
                    if image.format not in {"PNG", "JPEG"} or image.width * image.height > 25_000_000:
                        raise ValueError("invalid or oversized raster image")
                    image.verify()
                assets[str(path.relative_to(Path(asset_root).resolve()))] = digest(path)
                mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
                value[2][0] = "data:" + mime + ";base64," + base64.b64encode(path.read_bytes()).decode("ascii")
            if kind == "Link":
                target = value[2][0]
                # External citations may remain clickable, but are never fetched here.
                if not (target.startswith("#") or urlsplit(target).scheme in {"https", "http"}):
                    raise ValueError("file/local/active hyperlinks are not supported")
            if kind == "Math":
                math_count += 1
                if value[0]["t"] == "DisplayMath":
                    match = TAG.search(value[1])
                    labels.append(match[1] if match else None)
                    if match:
                        value[1] = value[1][:match.start()].rstrip()
                if "\\tag" in value[1]:
                    raise ValueError("equation tag must be a terminal numeric display tag")
            # This empty fenced div is the sole author-controlled page-break syntax.
            if kind == "Div" and "page-break" in value[0][1]:
                if value[1] or value[0][2]:
                    raise ValueError("page-break div must be empty")
                node.clear()
                node.update(t="RawBlock", c=["openxml", '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'])
                return
            for item in node.values():
                visit(item)
        elif isinstance(node, list):
            for item in node:
                visit(item)
    visit(ast)
    numbered = [label for label in labels if label is not None]
    if len(numbered) != len(set(numbered)):
        raise ValueError("duplicate equation numbers")
    return ast, labels, math_count, assets


def load_spec(path):
    spec = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    required = {"paper", "margin_cm", "body_pt", "line_spacing", "latin_font", "cjk_font",
                "heading_cjk_font", "math_font", "first_heading_is_title", "render_dpi", "timeout_seconds"}
    if set(spec) != required or spec["paper"] != "A4":
        raise ValueError("unsupported export configuration")
    bounds = {"margin_cm": (1.5, 4), "body_pt": (9, 14), "line_spacing": (1, 2),
              "render_dpi": (96, 200), "timeout_seconds": (10, 600)}
    for key, (low, high) in bounds.items():
        if isinstance(spec[key], bool) or not isinstance(spec[key], (int, float)) or not low <= spec[key] <= high:
            raise ValueError(f"out-of-range configuration: {key}")
    if not isinstance(spec["first_heading_is_title"], bool):
        raise ValueError("first_heading_is_title must be boolean")
    for key in ("latin_font", "cjk_font", "heading_cjk_font", "math_font"):
        if not isinstance(spec[key], str) or not spec[key].strip():
            raise ValueError(f"invalid font name: {key}")
    return spec


def format_docx(path, spec, labels, expected_math):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    document = Document(path)
    width = 21 - 2 * spec["margin_cm"]
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(spec["margin_cm"])
        section.left_margin = section.right_margin = Cm(spec["margin_cm"])
        section.footer_distance = Cm(1.25)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    for style in document.styles:
        if style.type not in (1, 2):
            continue
        style.font.name = "Consolas" if style.name.endswith("Tok") else spec["latin_font"]
        style.font.size = Pt(spec["body_pt"])
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), spec["heading_cjk_font"] if style.name.startswith("Heading") else spec["cjk_font"])
    for name in ("Source Code", "Verbatim Char"):
        if name in document.styles:
            document.styles[name].font.name = "Consolas"
    for name in ("Normal", "Body Text", "First Paragraph"):
        style = document.styles[name]
        style.paragraph_format.line_spacing = spec["line_spacing"]
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.widow_control = True
    for level, size in ((1, 15), (2, 13), (3, 12)):
        style = document.styles[f"Heading {level}"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    title_style = document.styles["Title"]
    title_style.font.size, title_style.font.bold = Pt(18), True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.keep_with_next = True
    if spec["first_heading_is_title"]:
        first = next((p for p in document.paragraphs if p.text.strip()), None)
        if first and first.style.name == "Heading 1":
            first.style = title_style
    for field in ("author", "last_modified_by", "comments", "identifier", "subject", "keywords", "category"):
        setattr(document.core_properties, field, "")
    math_pr = document.settings.element.find(qn("m:mathPr"))
    if math_pr is None:
        math_pr = OxmlElement("m:mathPr")
        document.settings.element.append(math_pr)
    font = math_pr.find(qn("m:mathFont"))
    if font is None:
        font = OxmlElement("m:mathFont")
        math_pr.append(font)
    font.set(qn("m:val"), spec["math_font"])
    # Standard data tables: repeated headers, no split rows. Never flatten runs/math.
    for table in document.tables:
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom", "insideH"):
            border = OxmlElement("w:" + edge)
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            borders.append(border)
        table._tbl.tblPr.append(borders)
        table.autofit = False
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.1
    displays = document.element.xpath(".//m:oMathPara")
    if len(displays) != len(labels):
        raise ValueError("Pandoc display-equation count mismatch")
    for math, label in zip(displays, labels):
        if label is None:
            continue
        paragraph = math.getparent()
        if paragraph.tag != qn("w:p") or paragraph.getparent().tag != qn("w:body"):
            raise ValueError("numbered display math must be a standalone body paragraph")
        if paragraph.xpath(".//w:t") or len(paragraph.xpath(".//m:oMathPara")) != 1:
            raise ValueError("numbered display math cannot share a paragraph with prose")
        table = document.add_table(rows=1, cols=3)
        table.autofit = False
        for col, cell, size in zip(table.columns, table.rows[0].cells, (1.0, width - 2, 1.0)):
            col.width = cell.width = Cm(size)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        center = table.cell(0, 1).paragraphs[0]
        center.alignment = WD_ALIGN_PARAGRAPH.CENTER
        center._p.append(math)
        right = table.cell(0, 2).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.add_run(f"({label})")
        for cell in table.rows[0].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph.addprevious(table._tbl)
        paragraph.getparent().remove(paragraph)
    if len(document.element.xpath(".//m:oMath")) != expected_math:
        raise ValueError("native equation count mismatch")
    for shape in document.inline_shapes:
        if shape.width > Cm(width):
            shape.height = int(shape.height * Cm(width) / shape.width)
            shape.width = Cm(width)
        paragraph = shape._inline
        while paragraph is not None and paragraph.tag != qn("w:p"):
            paragraph = paragraph.getparent()
        if paragraph is None:
            raise ValueError("image has no containing paragraph")
        props = paragraph.find(qn("w:pPr"))
        if props is None:
            props = OxmlElement("w:pPr")
            paragraph.insert(0, props)
        props.append(OxmlElement("w:keepNext"))
        alignment = props.find(qn("w:jc"))
        if alignment is None:
            alignment = OxmlElement("w:jc")
            props.append(alignment)
        alignment.set(qn("w:val"), "center")
    for properties in document.element.xpath(".//pic:cNvPr"):
        if properties.get("descr", "").startswith("data:"):
            properties.set("descr", "Embedded local raster")
    document.save(path)
    return {"native_math_count": expected_math, "display_math_count": len(labels),
            "visible_equation_numbers": [x for x in labels if x is not None],
            "data_tables": len(document.tables) - sum(x is not None for x in labels)}


def build(args):
    from pypdf import PdfReader
    source = inside(args.asset_root, args.md)
    if not source.is_file() or source.suffix.lower() != ".md":
        raise ValueError("source must be an existing Markdown file inside asset root")
    output = Path(args.out).resolve()
    if output.suffix.lower() != ".docx":
        raise ValueError("output must end in .docx")
    spec_path = Path(args.spec).resolve()
    spec = load_spec(spec_path)
    timeout = spec["timeout_seconds"]
    pandoc = executable(args.pandoc)
    soffice = executable(args.soffice) if args.pdf else None
    poppler = executable(args.pdftoppm) if args.pdf else None
    targets = [output, output.with_suffix(".export.json")]
    if args.pdf:
        targets += [output.with_suffix(".pdf"), output.with_suffix(".pages")]
    if any(target.exists() for target in targets):
        raise FileExistsError("output bundle exists; choose a new versioned output filename")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="paper-export-", dir=output.parent) as temp:
        temp = Path(temp)
        raw = run([pandoc, "--sandbox", "--from", READER, "--to", "json"],
                  data=source.read_bytes(), timeout=timeout, cwd=temp)
        ast, labels, count, assets = prepare_ast(json.loads(raw), source, args.asset_root)
        draft = temp / output.name
        command = [pandoc, "--sandbox", "--from", "json", "--to", "docx", "--standalone",
                   "--fail-if-warnings", "--output", str(draft)]
        if args.toc:
            command.append("--toc")
        run(command, data=json.dumps(ast).encode("utf-8"), timeout=timeout, cwd=temp)
        structure = format_docx(draft, spec, labels, count)
        versions = {"pandoc": run([pandoc, "--version"], timeout=timeout).decode("utf-8", "replace").splitlines()[0]}
        files = [(draft, output)]
        pages = []
        if args.pdf:
            pdf_dir = temp / "pdf"
            pdf_dir.mkdir()
            profile = (temp / "lo-profile").as_uri()
            run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                 "--outdir", str(pdf_dir), str(draft)], timeout=timeout, cwd=temp)
            pdf = pdf_dir / output.with_suffix(".pdf").name
            if not pdf.is_file() or not pdf.stat().st_size:
                raise RuntimeError("LibreOffice returned without a fresh PDF")
            reader = PdfReader(pdf)
            if not reader.pages:
                raise ValueError("empty PDF")
            page_dir = temp / "pages"
            page_dir.mkdir()
            run([poppler, "-png", "-r", str(spec["render_dpi"]), str(pdf), str(page_dir / "page")],
                timeout=timeout, cwd=temp)
            images = sorted(page_dir.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
            if len(images) != len(reader.pages):
                raise ValueError("PDF/render page count mismatch")
            for index, image in enumerate(images, 1):
                pages.append({"page": index, "file": f"{output.stem}.pages/{image.name}",
                              "sha256": digest(image), "status": "pending"})
            files.append((pdf, output.with_suffix(".pdf")))
            versions["libreoffice"] = run([soffice, "--version"], timeout=timeout).decode("utf-8", "replace").strip()
            files.append((page_dir, output.with_suffix(".pages")))
        artifacts = [{"file": target.name, "sha256": digest(path)} for path, target in files if path.is_file()]
        report = {"schema_version": 1, "build_id": str(uuid.uuid4()), "status": "pending_visual_review",
                  "submission_ready": False, "source_sha256": digest(source), "spec_sha256": digest(spec_path),
                  "exporter_sha256": digest(__file__), "assets": assets, "structure": structure,
                  "tools": versions, "artifacts": artifacts, "page_count": len(pages), "pages": pages,
                  "review_scope": "AI page review is not human acceptance, official-rule or submission approval"}
        report_path = temp / "report.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        # Manifest is published last. A partial bundle after an OS failure is not accepted.
        for path, target in files:
            if target.exists():
                raise FileExistsError(target)
            path.rename(target)
        report_path.rename(output.with_suffix(".export.json"))
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--md", type=Path, required=True)
    parser.add_argument("--asset-root", type=Path, required=True)
    parser.add_argument("--spec", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--pdf", action="store_true", help="Export PDF and render every page; otherwise DOCX only")
    parser.add_argument("--toc", action="store_true", help="Opt in only after rules/template review")
    parser.add_argument("--pandoc", default="pandoc")
    parser.add_argument("--soffice", default="soffice")
    parser.add_argument("--pdftoppm", default="pdftoppm")
    args = parser.parse_args()
    try:
        report = build(args)
    except ImportError as exc:
        parser.exit(2, f"missing or broken Python dependency: {exc}; run doctor.py with this interpreter\n")
    except (OSError, ValueError, RuntimeError, subprocess.SubprocessError) as exc:
        parser.exit(2, f"export failed: {exc}\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

