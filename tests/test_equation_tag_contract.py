"""Synthetic regression for the public exporter's existing tag contract."""
import importlib.util
from pathlib import Path
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('public_export',ROOT/'tools/paper-export/export_paper.py')
export=importlib.util.module_from_spec(spec)
spec.loader.exec_module(export)


class EquationTagContract(unittest.TestCase):
    def prepare(self, items):
        ast={'blocks':[{'t':'Para','c':[{'t':'Math','c':[{'t':kind},value]}]} for kind,value in items]}
        with tempfile.TemporaryDirectory() as d:
            return export.prepare_ast(ast,Path(d)/'probe.md',Path(d))

    def test_mixed_equations(self):
        ast,labels,count,_=self.prepare([('DisplayMath',r'x=1\tag{4-1}'),('DisplayMath','y=2'),('InlineMath','z^2')])
        self.assertEqual(labels,['4-1',None])
        self.assertEqual(count,3)
        self.assertEqual(ast['blocks'][0]['c'][0]['c'][1],'x=1')
        self.assertEqual(ast['blocks'][2]['c'][0]['c'][1],'z^2')

    def test_duplicate_rejected(self):
        with self.assertRaisesRegex(ValueError,'duplicate'):
            self.prepare([('DisplayMath',r'x\tag{1}'),('DisplayMath',r'y\tag{1}')])

    def test_invalid_tags_rejected(self):
        for kind,text in [('InlineMath',r'x\tag{1}'),('DisplayMath',r'x\tag{A}'),('DisplayMath',r'x\tag{1}y')]:
            with self.subTest(text=text),self.assertRaises(ValueError):self.prepare([(kind,text)])

    def test_native_math_and_visible_number_after_formatting(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.enum.style import WD_STYLE_TYPE
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'probe.docx';doc=Document()
            doc.styles.add_style('First Paragraph',WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph('Reference (4-1)')
            for value in ('x=1','y=2'):
                p=doc.add_paragraph(); display=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath')
                run=OxmlElement('m:r');text=OxmlElement('m:t');text.text=value
                run.append(text);math.append(run);display.append(math);p._p.append(display)
            doc.save(path)
            export.format_docx(path,export.load_spec(ROOT/'doc-export-specs/paper-default.json'),['4-1',None],2)
            doc=Document(path)
            self.assertEqual(len(doc.element.xpath('.//m:oMath')),2)
            self.assertEqual(len(doc.tables),1)
            self.assertEqual(doc.tables[0].cell(0,2).text,'(4-1)')
            self.assertEqual(doc.paragraphs[0].text,'Reference (4-1)')


if __name__=='__main__':unittest.main()
